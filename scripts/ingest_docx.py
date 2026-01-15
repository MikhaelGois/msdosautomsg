"""INGEST DOCX -> JSON catalog

Usage:
  python scripts/ingest_docx.py path/to/templates.docx --out data/templates_catalog.json

This script extracts template sections from a DOCX file and writes a JSON array with fields:
  - id (generated)
  - name
  - subject
  - body (plain text)
  - slots (list of detected placeholders)

It looks for obvious template boundaries (Heading styles) and for an explicit "Subject:" line.
Requires: python-docx
  pip install python-docx
"""
from __future__ import annotations
import re
import json
import sys
from pathlib import Path
from docx import Document
from typing import List, Dict


PLACEHOLDER_RE = re.compile(r"{{\s*([^}]+)\s*}}|\[([^\]]+)\]|\{([^}]+)\}")


def extract_placeholders(text: str) -> List[str]:
    matches = PLACEHOLDER_RE.findall(text)
    slots = []
    for a, b, c in matches:
        v = a or b or c
        v = v.strip()
        if v and v not in slots:
            slots.append(v)
    return slots


def ingest_docx(path: Path) -> List[Dict]:
    doc = Document(path)
    templates = []
    cur = None

    def push_current():
        if cur and (cur.get("body") or cur.get("subject")):
            body = "\n".join(cur.get("body", []))
            slots = extract_placeholders(body)
            # also scan subject
            if cur.get("subject"):
                slots += [s for s in extract_placeholders(cur["subject"]) if s not in slots]
            templates.append({
                "id": len(templates) + 1,
                "name": cur.get("name") or f"template_{len(templates)+1}",
                "subject": cur.get("subject") or "",
                "body": body,
                "slots": slots,
            })

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            # preserve blank lines in body
            if cur is not None:
                cur.setdefault("body", []).append("")
            continue

        style_name = (p.style.name or "").lower()
        # treat headings as new template boundary
        if style_name.startswith("heading") or style_name in ("title",):
            push_current()
            cur = {"name": text, "body": []}
            continue

        # explicit Subject: line
        if text.lower().startswith("subject:"):
            if cur is None:
                cur = {"name": text.replace("Subject:", "").strip(), "body": []}
            cur["subject"] = text.split(":", 1)[1].strip()
            continue

        # also detect a pattern like "Template: <name>"
        if text.lower().startswith("template:"):
            push_current()
            cur = {"name": text.split(":", 1)[1].strip(), "body": []}
            continue

        # otherwise append to current body (or start unnamed template)
        if cur is None:
            cur = {"name": None, "body": [text]}
        else:
            cur.setdefault("body", []).append(text)

    push_current()
    return templates


def main(argv: List[str] | None = None) -> int:
    argv = argv or sys.argv[1:]
    if not argv:
        print("Usage: python scripts/ingest_docx.py path/to/file.docx [--out output.json]")
        return 1
    path = Path(argv[0])
    out = Path("data/templates_catalog.json")
    if "--out" in argv:
        i = argv.index("--out")
        if i + 1 < len(argv):
            out = Path(argv[i + 1])

    if not path.exists():
        print(f"File not found: {path}")
        return 2

    templates = ingest_docx(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    with open(out, "w", encoding="utf-8") as f:
        json.dump(templates, f, ensure_ascii=False, indent=2)

    print(f"Wrote {len(templates)} templates to {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
